#!/usr/bin/env python3
"""Normalize a Power BI Retail Sales Excel export without mixing hierarchy levels."""

from __future__ import annotations

import argparse
import json
import math
import re
from datetime import date
from pathlib import Path

from openpyxl import load_workbook


EXPECTED = {
    "Department",
    "SubDepartment",
    "Class",
    "SubClass",
    "ItemId-Desc",
    "Net Sales$",
    "Net Margin$",
    "Net Margin%",
    "Movement",
    "LY Net Sales$",
    "YoY Movement",
}


def args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", required=True)
    parser.add_argument("--json", required=True)
    parser.add_argument("--start", required=True)
    parser.add_argument("--end", required=True)
    parser.add_argument("--comparison-start", required=True)
    parser.add_argument("--comparison-end", required=True)
    parser.add_argument("--store", default="Pet Supplies Plus #212 — Battle Creek, MI")
    parser.add_argument("--source-view", default="Power BI Retail Sales — Item Summary")
    return parser.parse_args()


def finite(value, default=0.0):
    if value is None or value == "":
        return default
    try:
        number = float(value)
    except (TypeError, ValueError):
        return default
    return number if math.isfinite(number) else default


def clean(value):
    return "" if value is None else str(value).strip()


def row_kind(row):
    department, subdepartment, class_name, subclass, item = row[:5]
    if clean(department).lower() == "total":
        return "grand_total"
    if clean(subdepartment).lower() == "total":
        return "department"
    if clean(class_name).lower() == "total":
        return "subdepartment"
    if clean(subclass).lower() == "total":
        return "class"
    if clean(item).lower() == "total":
        return "subclass"
    if clean(item):
        return "product"
    return "other"


def split_item(value):
    text = clean(value)
    match = re.match(r"^([^\s]+)\s+-\s+(.*)$", text)
    if match:
        return match.group(1), match.group(2)
    return text, text


def main():
    options = args()
    for raw in (options.start, options.end, options.comparison_start, options.comparison_end):
        date.fromisoformat(raw)

    workbook = load_workbook(options.input, read_only=True, data_only=True)
    selected = None
    headers = None
    for sheet in workbook.worksheets:
        first = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
        candidate = [clean(value) for value in first]
        if EXPECTED.issubset(set(candidate)):
            selected = sheet
            headers = candidate
            break
    if selected is None:
        raise SystemExit("No worksheet contains the expected Power BI Retail Sales fields")

    index = {name: position for position, name in enumerate(headers)}
    aliases = {
        "department": "Department",
        "subdepartment": "SubDepartment",
        "class": "Class",
        "subclass": "SubClass",
        "item": "ItemId-Desc",
        "sales": "Net Sales$",
        "cost": "Cost$",
        "scanback": "Scanback$",
        "margin": "Net Margin$",
        "margin_pct": "Net Margin%",
        "units": "Movement",
        "comparison_sales": "LY Net Sales$",
        "sales_pct_change": "YoY Net Sales$ %",
        "margin_rate_change": "YoY Net Margin%",
        "unit_change": "YoY Movement",
        "avg_unit_retail": "Average Unit Net Retail$",
    }

    records = []
    filters = ""
    excluded = {"blank_or_note": 0}
    for values in selected.iter_rows(min_row=2, values_only=True):
        padded = list(values) + [None] * max(0, len(headers) - len(values))
        if all(value in (None, "") for value in padded):
            excluded["blank_or_note"] += 1
            continue
        if clean(padded[0]).startswith("Applied filters:"):
            filters = clean(padded[0])
            excluded["blank_or_note"] += 1
            continue
        source = {key: padded[index[column]] if column in index else None for key, column in aliases.items()}
        kind = row_kind([
            source["department"], source["subdepartment"], source["class"],
            source["subclass"], source["item"],
        ])
        if kind == "other":
            excluded["blank_or_note"] += 1
            continue

        current_sales = finite(source["sales"])
        comparison_sales = finite(source["comparison_sales"])
        current_units = finite(source["units"])
        unit_change = finite(source["unit_change"])
        comparison_units = current_units - unit_change
        current_margin = finite(source["margin"])
        current_margin_pct = finite(source["margin_pct"])
        margin_rate_change = finite(source["margin_rate_change"])
        comparison_margin_pct = current_margin_pct - margin_rate_change
        comparison_margin = comparison_sales * comparison_margin_pct
        item_id, description = split_item(source["item"])

        records.append({
            "kind": kind,
            "department": clean(source["department"]),
            "subdepartment": clean(source["subdepartment"]),
            "class": clean(source["class"]),
            "subclass": clean(source["subclass"]),
            "itemId": item_id if kind == "product" else "",
            "description": description if kind == "product" else "",
            "currentSales": current_sales,
            "comparisonSales": comparison_sales,
            "currentUnits": current_units,
            "comparisonUnits": comparison_units,
            "currentMargin": current_margin,
            "comparisonMargin": comparison_margin,
            "currentMarginPct": current_margin_pct,
            "comparisonMarginPct": comparison_margin_pct,
            "scanback": finite(source["scanback"]),
            "currentAvgUnitRetail": finite(source["avg_unit_retail"]),
            "comparisonAvgUnitRetail": comparison_sales / comparison_units if comparison_units else 0.0,
            "sourceSalesPctChange": finite(source["sales_pct_change"]),
            "sourceUnitChange": unit_change,
            "sourceMarginRateChange": margin_rate_change,
        })

    counts = {}
    for record in records:
        counts[record["kind"]] = counts.get(record["kind"], 0) + 1
    if counts.get("grand_total") != 1:
        raise SystemExit(f"Expected exactly one grand total; found {counts.get('grand_total', 0)}")

    payload = {
        "metadata": {
            "store": options.store,
            "sourceFile": str(Path(options.input).resolve()),
            "sourceView": options.source_view,
            "sheet": selected.title,
            "currentStart": options.start,
            "currentEnd": options.end,
            "comparisonStart": options.comparison_start,
            "comparisonEnd": options.comparison_end,
            "comparisonMethod": "Power BI LY measures; fiscal-aligned prior-year period",
            "appliedFilters": filters,
            "recordCounts": counts,
            "excluded": excluded,
        },
        "records": records,
    }
    destination = Path(options.json)
    destination.parent.mkdir(parents=True, exist_ok=True)
    destination.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"json": str(destination.resolve()), "counts": counts}, ensure_ascii=False))


if __name__ == "__main__":
    main()
